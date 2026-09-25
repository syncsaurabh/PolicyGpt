from datetime import date, datetime, timedelta, timezone
import io
import json
import logging
from typing import Any, Dict, List, Optional, Tuple, cast
from fastapi import HTTPException, status
from sqlalchemy import desc
from sqlalchemy.orm import Session
from app.models.policy import Policy
from app.models.report import Report, ReportFormat, ReportType
from app.models.scheme import Scheme
from app.models.user import User
from app.models.user_activity import UserActivity
from app.schemas.report import ReportFilterParams

logger = logging.getLogger(__name__)


class ReportService:
    # --- Data Extraction Helpers ---

    @staticmethod
    def get_policy_report_data(
        db: Session,
        filters: Optional[ReportFilterParams] = None,
    ) -> List[Dict[str, Any]]:
        """Fetch policy dataset based on filters."""
        query = db.query(Policy)
        if filters:
            if filters.start_date:
                query = query.filter(Policy.created_at >= filters.start_date)
            if filters.end_date:
                query = query.filter(Policy.created_at < filters.end_date + timedelta(days=1))
            if filters.department:
                query = query.filter(Policy.department.ilike(f"%{filters.department.strip()}%"))
            if filters.category:
                query = query.filter(Policy.category.ilike(f"%{filters.category.strip()}%"))
            if filters.state:
                query = query.filter(Policy.state.ilike(f"%{filters.state.strip()}%"))
            if filters.status:
                query = query.filter(Policy.status.ilike(filters.status.strip()))

        policies = query.order_by(desc(Policy.created_at)).all()
        return [
            {
                "id": p.id,
                "title": p.title,
                "category": p.category or "N/A",
                "department": p.department or "N/A",
                "state": p.state or "National",
                "status": p.status,
                "created_at": p.created_at.strftime("%Y-%m-%d") if p.created_at else "N/A",
            }
            for p in policies
        ]

    @staticmethod
    def get_scheme_report_data(
        db: Session,
        filters: Optional[ReportFilterParams] = None,
    ) -> List[Dict[str, Any]]:
        """Fetch scheme dataset based on filters."""
        query = db.query(Scheme)
        if filters:
            if filters.start_date:
                query = query.filter(Scheme.created_at >= filters.start_date)
            if filters.end_date:
                query = query.filter(Scheme.created_at < filters.end_date + timedelta(days=1))
            if filters.department:
                query = query.filter(Scheme.department.ilike(f"%{filters.department.strip()}%"))
            if filters.category:
                query = query.filter(Scheme.category.ilike(f"%{filters.category.strip()}%"))
            if filters.state:
                query = query.filter(Scheme.state.ilike(f"%{filters.state.strip()}%"))
            if filters.status:
                query = query.filter(Scheme.status.ilike(filters.status.strip()))

        schemes = query.order_by(desc(Scheme.created_at)).all()
        return [
            {
                "id": s.id,
                "name": s.name,
                "category": s.category or "N/A",
                "department": s.department or "N/A",
                "target_audience": s.target_audience or "General",
                "status": s.status,
                "created_at": s.created_at.strftime("%Y-%m-%d") if s.created_at else "N/A",
            }
            for s in schemes
        ]

    @staticmethod
    def get_department_report_data(
        db: Session,
        filters: Optional[ReportFilterParams] = None,
    ) -> List[Dict[str, Any]]:
        """Aggregate department reporting data."""
        p_depts = db.query(Policy.department).filter(Policy.department.isnot(None))
        s_depts = db.query(Scheme.department).filter(Scheme.department.isnot(None))

        if filters and filters.department:
            p_depts = p_depts.filter(Policy.department.ilike(f"%{filters.department.strip()}%"))
            s_depts = s_depts.filter(Scheme.department.ilike(f"%{filters.department.strip()}%"))

        dept_names = set([r[0].strip() for r in p_depts.distinct().all() if r[0]])
        dept_names.update([r[0].strip() for r in s_depts.distinct().all() if r[0]])

        data: List[Dict[str, Any]] = []
        for dept in sorted(dept_names):
            p_count = db.query(Policy).filter(Policy.department.ilike(dept)).count()
            s_count = db.query(Scheme).filter(Scheme.department.ilike(dept)).count()
            p_pub = db.query(Policy).filter(Policy.department.ilike(dept), Policy.status.in_(["PUBLISHED", "APPROVED"])).count()
            s_act = db.query(Scheme).filter(Scheme.department.ilike(dept), Scheme.status.in_(["ACTIVE", "PUBLISHED"])).count()
            data.append({
                "department": dept,
                "total_policies": p_count,
                "published_policies": p_pub,
                "total_schemes": s_count,
                "active_schemes": s_act,
            })
        return data

    @staticmethod
    def get_user_activity_report_data(
        db: Session,
        filters: Optional[ReportFilterParams] = None,
    ) -> List[Dict[str, Any]]:
        """Fetch user activity logs dataset."""
        query = db.query(UserActivity)
        if filters:
            if filters.start_date:
                query = query.filter(UserActivity.created_at >= filters.start_date)
            if filters.end_date:
                query = query.filter(UserActivity.created_at < filters.end_date + timedelta(days=1))
            if filters.event_type:
                query = query.filter(UserActivity.event_type == filters.event_type)

        activities = query.order_by(desc(UserActivity.created_at)).limit(500).all()
        return [
            {
                "id": a.id,
                "user_id": a.user_id or "Guest",
                "event_type": a.event_type,
                "resource_type": a.resource_type or "N/A",
                "resource_id": a.resource_id or "N/A",
                "ip_address": a.ip_address or "N/A",
                "created_at": a.created_at.strftime("%Y-%m-%d %H:%M:%S") if a.created_at else "N/A",
            }
            for a in activities
        ]

    # --- PDF Generation Engine (ReportLab) ---

    @staticmethod
    def export_pdf(
        title: str,
        headers: List[str],
        data_rows: List[List[Any]],
        filters_summary: Dict[str, Any],
    ) -> io.BytesIO:
        """Generate a professionally formatted PDF document stream."""
        try:
            from reportlab.lib import colors  # type: ignore
            from reportlab.lib.pagesizes import A4, landscape  # type: ignore
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # type: ignore
        except ImportError:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="PDF generation engine ('reportlab') is not installed.",
            )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1A365D"),
            alignment=1,  # Centered
        )
        subtitle_style = ParagraphStyle(
            "ReportSubtitle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#4A5568"),
            alignment=1,
        )
        filter_style = ParagraphStyle(
            "FilterText",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#2D3748"),
        )
        cell_style = ParagraphStyle(
            "CellText",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#1A202C"),
        )
        header_cell_style = ParagraphStyle(
            "HeaderCellText",
            parent=styles["Normal"],
            fontSize=9,
            leading=11,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )

        story: List[Any] = []

        # Document Header
        story.append(Paragraph(title, title_style))
        gen_date = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        story.append(Paragraph(f"PolicyGPT Platform Official Report | Generated at: {gen_date}", subtitle_style))
        story.append(Spacer(1, 10))

        # Filters Applied block
        if filters_summary:
            f_text = " | ".join([f"<b>{k}:</b> {v}" for k, v in filters_summary.items() if v])
            if f_text:
                story.append(Paragraph(f"<b>Applied Filters:</b> {f_text}", filter_style))
                story.append(Spacer(1, 8))

        # Table data construction
        table_content: List[List[Any]] = []
        # Header row
        table_content.append([Paragraph(h, header_cell_style) for h in headers])

        # Data rows
        for row in data_rows:
            formatted_row = []
            for item in row:
                text_val = str(item) if item is not None else ""
                formatted_row.append(Paragraph(text_val, cell_style))
            table_content.append(formatted_row)

        if not data_rows:
            empty_row = [Paragraph("No records found matching criteria", cell_style)] + [Paragraph("", cell_style)] * (len(headers) - 1)
            table_content.append(empty_row)

        table = Table(table_content, repeatRows=1)
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B6CB0")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F7FAFC")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ])
        )
        story.append(table)
        story.append(Spacer(1, 12))
        story.append(Paragraph(f"Total Records: {len(data_rows)}", subtitle_style))

        doc.build(story)
        buffer.seek(0)
        return buffer

    # --- Excel Generation Engine (OpenPyXL) ---

    @staticmethod
    def export_excel(
        sheet_title: str,
        title: str,
        headers: List[str],
        data_rows: List[List[Any]],
        filters_summary: Dict[str, Any],
    ) -> io.BytesIO:
        """Generate a professionally styled Excel spreadsheet stream."""
        try:
            from openpyxl import Workbook  # type: ignore
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore
        except ImportError:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Excel generation engine ('openpyxl') is not installed.",
            )

        wb = Workbook()
        ws: Any = wb.active if wb.active is not None else wb.create_sheet()
        ws.title = sheet_title[:31]  # Excel tab name max 31 chars

        # Style definitions
        title_font = Font(name="Calibri", size=14, bold=True, color="1A365D")
        meta_font = Font(name="Calibri", size=10, italic=True, color="4A5568")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
        alt_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
        border_side = Side(style="thin", color="E2E8F0")
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        align_left = Alignment(horizontal="left", vertical="center")

        # Row 1: Title
        ws.append([title])
        ws.cell(row=1, column=1).font = title_font

        # Row 2: Metadata
        gen_str = f"Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | PolicyGPT Platform"
        ws.append([gen_str])
        ws.cell(row=2, column=1).font = meta_font

        # Row 3: Filters
        if filters_summary:
            f_str = "Filters: " + ", ".join([f"{k}={v}" for k, v in filters_summary.items() if v])
            ws.append([f_str])
            ws.cell(row=3, column=1).font = meta_font
            current_row = 5
        else:
            current_row = 4

        # Header Row
        ws.cell(row=current_row - 1, column=1).value = ""  # Spacer
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_left
            cell.border = thin_border

        # Data Rows
        for r_idx, row in enumerate(data_rows, start=current_row + 1):
            ws.append(row)
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.alignment = align_left
                cell.border = thin_border
                if (r_idx - current_row) % 2 == 0:
                    cell.fill = alt_fill

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # --- Report Log Persistence ---

    @staticmethod
    def log_generated_report(
        db: Session,
        title: str,
        report_type: str,
        file_format: str,
        record_count: int,
        generated_by: Optional[int] = None,
        filters: Optional[ReportFilterParams] = None,
    ) -> Report:
        """Record report generation event into database for history and audit tracking."""
        params_str = None
        if filters:
            params_str = json.dumps(filters.model_dump(exclude_unset=True), default=str)

        report = Report(
            generated_by=generated_by,
            title=title,
            report_type=report_type,
            file_format=file_format,
            record_count=record_count,
            parameters_json=params_str,
            status="COMPLETED",
        )
        db.add(report)
        db.commit()
        db.refresh(report)
        return report

    @staticmethod
    def list_reports(
        db: Session,
        page: int = 1,
        page_size: int = 10,
        report_type: Optional[str] = None,
    ) -> Tuple[List[Report], int, int]:
        """List previously generated report logs."""
        query = db.query(Report)
        if report_type:
            query = query.filter(Report.report_type == report_type)

        total_count = query.count()
        total_pages = (total_count + page_size - 1) // page_size if total_count > 0 else 1
        offset = (page - 1) * page_size
        results = query.order_by(desc(Report.created_at)).offset(offset).limit(page_size).all()

        return results, total_count, total_pages
